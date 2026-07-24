"""
streamlit_prueba.py — Interfaz Streamlit del Asistente de Memoria Anual
Departamento de Desarrollo Local y Empleo · Ayuntamiento de San Sebastián de los Reyes

Basado en el diseño original de Paloma (app.py), reconectado al pipeline
de agentes actual (src_agents.graph.workflow), en vez de a los modulos
antiguos (redactor.py / validation/revisor.py / rag/loader.py).

Se quitaron los selectores de modelo, temperatura y secciones: el pipeline
nuevo no expone esos parametros (decision ya tomada: el personal municipal
no debe elegir modelo ni temperatura, y el pipeline siempre genera el
informe completo en espanol + ingles, no seccion a seccion).
"""

import base64
import io
import sys
import tempfile
import time
from pathlib import Path

import streamlit as st
from docx import Document as DocxDocument

from dotenv import load_dotenv
load_dotenv()

# -- Ruta raíz al path para importar src_agents.*
ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from src_agents.graph.workflow import pipeline

# ---------------------------------------------------------------------------
# Helper — construir .docx en memoria (español + inglés + incidencias)
# ---------------------------------------------------------------------------
def _construir_docx(draft_es: str, draft_en: str, incidencias: list[str]) -> bytes:
    doc = DocxDocument()
    doc.add_heading("Memoria Anual de Actividades", level=0)
    doc.add_paragraph(
        "Departamento de Desarrollo Local y Empleo\n"
        "Ayuntamiento de San Sebastián de los Reyes\n"
        "Generado por el Asistente IA · Factoría F5"
    )
    doc.add_page_break()

    doc.add_heading("Informe (Español)", level=1)
    doc.add_paragraph(draft_es)

    if draft_en:
        doc.add_page_break()
        doc.add_heading("Report (English)", level=1)
        doc.add_paragraph(draft_en)

    if incidencias:
        doc.add_page_break()
        doc.add_heading("Notas de validación (revisión humana)", level=1)
        nota = doc.add_paragraph()
        nota.add_run(
            "Estos datos no se han encontrado literalmente en el texto generado. "
            "No bloquea el informe — queda para revisión humana:"
        ).italic = True
        for inc in incidencias:
            doc.add_paragraph(inc, style="List Bullet")

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


# ---------------------------------------------------------------------------
# Configuración de página
# ---------------------------------------------------------------------------
st.set_page_config(
    page_title="Asistente Memoria Anual · Sanse",
    page_icon="🏛️",
    layout="wide",
    initial_sidebar_state="expanded",
)

# ---------------------------------------------------------------------------
# Logo SVG → base64
# ---------------------------------------------------------------------------
_LOGO_PATH = Path(__file__).parent / "assets" / "logo_sanse.svg"
_logo_b64 = base64.b64encode(_LOGO_PATH.read_bytes()).decode() if _LOGO_PATH.exists() else ""
_LOGO_HTML = (
    f'<img src="data:image/svg+xml;base64,{_logo_b64}" '
    f'alt="Ayuntamiento de San Sebastián de los Reyes" style="height:56px;display:block;">'
    if _logo_b64 else '<span style="font-size:2rem">🏛️</span>'
)

# ---------------------------------------------------------------------------
# CSS — estética institucional Ayuntamiento de San Sebastián de los Reyes
# (sin cambios respecto al diseño original de Paloma)
# ---------------------------------------------------------------------------
st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Open+Sans:wght@400;600;700&display=swap');
:root {
    --sanse-red:    #A3132F;
    --sanse-red-dk: #7D0E22;
    --sanse-red-lt: #F9E8EB;
    --sanse-gray:   #F4F4F4;
    --sanse-border: #E0E0E0;
    --sanse-text:   #1A1A1A;
    --sanse-muted:  #5A5A5A;
    --sanse-ok:     #1A6B3C;
    --sanse-warn:   #8B6000;
    --radius:       6px;
}
html, body, .stApp {
    font-family: 'Open Sans', Arial, sans-serif;
    background: #FFFFFF;
    color: var(--sanse-text);
}
header[data-testid="stHeader"], #MainMenu, footer { display: none !important; }
.block-container { padding-top: 0 !important; }
.sanse-header {
    display: flex;
    align-items: center;
    gap: 1.4rem;
    background: #FFFFFF;
    border-bottom: 3px solid var(--sanse-red);
    padding: 1rem 1.8rem 0.9rem;
    margin-bottom: 1.8rem;
}
.sanse-header-text h1 { font-size:1.35rem; font-weight:700; color:var(--sanse-text); margin:0 0 2px; }
.sanse-header-text p  { font-size:.82rem; color:var(--sanse-muted); margin:0; }
section[data-testid="stSidebar"] { background: var(--sanse-gray) !important; border-right: 1px solid var(--sanse-border); }
section[data-testid="stSidebar"] * { color: var(--sanse-text) !important; }
section[data-testid="stSidebar"] h3 { color: var(--sanse-red) !important; font-weight:700; font-size:.95rem; text-transform:uppercase; letter-spacing:.05em; }
section[data-testid="stSidebar"] hr { border-color: var(--sanse-border) !important; }
.section-card { background:white; border:1px solid var(--sanse-border); border-top:4px solid var(--sanse-red); border-radius:var(--radius); padding:1.2rem 1.6rem; margin-bottom:1.2rem; box-shadow:0 1px 4px rgba(0,0,0,.06); }
.section-card h3 { color:var(--sanse-red); margin-top:0; font-size:1rem; }
.badge-ok   { background:var(--sanse-ok);   color:white; padding:2px 10px; border-radius:20px; font-size:.75rem; font-weight:600; }
.badge-warn { background:var(--sanse-warn);  color:white; padding:2px 10px; border-radius:20px; font-size:.75rem; font-weight:600; }
[data-testid="stFileUploaderDropzone"] { background:white !important; border:2px dashed var(--sanse-border) !important; border-radius:var(--radius) !important; }
[data-testid="stFileUploaderDropzone"] * { color:var(--sanse-text) !important; }
[data-testid="stFileUploaderDropzoneInstructions"] > div > span { visibility:hidden; display:block; height:0; }
[data-testid="stFileUploaderDropzoneInstructions"] > div > span::before { visibility:visible; display:block; height:auto; content:"Arrastra y suelta los archivos aquí"; font-weight:600; font-size:.95rem; color:var(--sanse-text) !important; }
[data-testid="stFileUploaderDropzoneInstructions"] > div > small { visibility:hidden; display:block; height:0; }
[data-testid="stFileUploaderDropzoneInstructions"] > div > small::before { visibility:visible; display:block; height:auto; content:"Límite 200 MB por archivo  •  PDF, DOCX, XLSX"; font-size:.82rem; color:var(--sanse-muted) !important; }
[data-testid="stFileUploaderDropzone"] button { background:var(--sanse-red) !important; color:white !important; border:none !important; border-radius:var(--radius) !important; font-weight:600 !important; font-size:0 !important; padding:.5rem 1.2rem !important; }
[data-testid="stFileUploaderDropzone"] button::before { content:"Seleccionar archivos"; font-size:.85rem; font-weight:600; color:white; }
[data-testid="stFileUploaderDropzone"] button:hover { background:var(--sanse-red-dk) !important; }
div.stButton > button[kind="primary"] { background:var(--sanse-red) !important; color:white !important; border:none !important; border-radius:3px !important; padding:.65rem 2.4rem !important; font-weight:700 !important; font-size:.95rem !important; letter-spacing:.08em !important; text-transform:uppercase !important; transition:background .2s, box-shadow .2s; box-shadow:0 2px 6px rgba(163,19,47,.30) !important; }
div.stButton > button[kind="primary"]:hover { background:var(--sanse-red-dk) !important; box-shadow:0 4px 12px rgba(163,19,47,.40) !important; }
.zona-titulo { font-size:1rem; font-weight:700; color:var(--sanse-text); border-left:4px solid var(--sanse-red); padding-left:.7rem; margin:1.4rem 0 .8rem; }
hr { border-color: var(--sanse-border) !important; }
div[data-testid="stInfo"] { background:var(--sanse-red-lt) !important; border-left:4px solid var(--sanse-red) !important; border-radius:var(--radius); color:var(--sanse-text) !important; }
@media (max-width: 768px) {
    .sanse-header { flex-direction:column; align-items:flex-start; gap:.8rem; padding:.8rem 1rem; }
    .sanse-header-text h1 { font-size:1.1rem; }
    .block-container { padding-left:.8rem !important; padding-right:.8rem !important; }
    [data-testid="column"] { width:100% !important; flex:1 1 100% !important; min-width:100% !important; }
    div.stButton > button[kind="primary"] { width:100% !important; }
    .section-card { padding:.9rem 1rem; }
}
</style>
""", unsafe_allow_html=True)

# ---------------------------------------------------------------------------
# Cabecera
# ---------------------------------------------------------------------------
st.markdown(f"""
<div class="sanse-header">
    {_LOGO_HTML}
    <div class="sanse-header-text">
        <h1>Asistente de Generación de Memoria Anual</h1>
        <p>Departamento de Desarrollo Local y Empleo &nbsp;·&nbsp; Ayuntamiento de San Sebastián de los Reyes</p>
    </div>
</div>
""", unsafe_allow_html=True)

# ---------------------------------------------------------------------------
# Sidebar — sin selectores de modelo/temperatura/secciones (no aplican al
# pipeline nuevo: el modelo es una decisión interna, no del usuario, y el
# pipeline siempre genera el informe completo en español + inglés).
# ---------------------------------------------------------------------------
with st.sidebar:
    st.markdown(_LOGO_HTML, unsafe_allow_html=True)
    st.markdown("<br>", unsafe_allow_html=True)
    st.markdown("### ℹ️ Cómo funciona")
    st.caption(
        "El sistema procesa tus documentos en varios pasos automáticos: "
        "lectura, interpretación de datos, redacción, revisión y traducción. "
        "El resultado se genera en español e inglés."
    )
    st.divider()
    st.caption("Proyecto pedagógico · Factoría F5 & Ayuntamiento Sanse · 2026")

# ---------------------------------------------------------------------------
# Zona principal — carga
# ---------------------------------------------------------------------------
col_upload, col_info = st.columns([2, 1])
with col_upload:
    st.markdown('<p class="zona-titulo">📂 Sube los documentos de entrada</p>', unsafe_allow_html=True)
    archivos_subidos = st.file_uploader(
        "Formatos admitidos: PDF, Word (.docx), Excel (.xlsx)",
        type=["pdf", "docx", "xlsx"],
        accept_multiple_files=True,
        help="Informes cuatrimestrales, convenios, memorias de la agencia de colocación.",
    )
with col_info:
    st.markdown('<p class="zona-titulo">📌 Documentos esperados</p>', unsafe_allow_html=True)
    st.info(
        "- Informe financiero cuatrimestral\n"
        "- Informes de seguimiento de convenios\n"
        "- Memoria agencia de colocación\n\n"
        "_Puedes subir varios archivos a la vez._"
    )

# ---------------------------------------------------------------------------
# Botón principal
# ---------------------------------------------------------------------------
st.divider()
generar_col, _ = st.columns([1, 3])
with generar_col:
    boton_generar = st.button(
        "Generar Memoria",
        type="primary",
        disabled=not archivos_subidos,
    )

# ---------------------------------------------------------------------------
# Pipeline de generación (nuevo: src_agents.graph.workflow.pipeline)
# ---------------------------------------------------------------------------
if boton_generar:
    # 1. Guardar los archivos subidos en una carpeta temporal — el agente
    #    Ingesta espera una carpeta (extractor_generico.extraer_carpeta),
    #    no una lista de rutas de archivo sueltas.
    with st.status("📄 Preparando documentos…", expanded=True) as status_carga:
        tmp_dir = Path(tempfile.mkdtemp(prefix="sanse_streamlit_"))
        for archivo in archivos_subidos:
            (tmp_dir / archivo.name).write_bytes(archivo.read())
            st.write(f"✅ **{archivo.name}** guardado para procesar")
        status_carga.update(label=f"✅ {len(archivos_subidos)} documento(s) listo(s)", state="complete")

    # 2. Ejecutar el pipeline completo
    with st.spinner("Generando el informe — esto puede tardar varios minutos…"):
        t0 = time.time()
        try:
            resultado = pipeline.invoke({"uploaded_files": [str(tmp_dir)]})
            elapsed = time.time() - t0
        except Exception as exc:
            st.error(f"Error generando el informe: {exc}")
            st.stop()

    draft_es = resultado.get("draft", "")
    draft_en = resultado.get("draft_en", "")
    review = resultado.get("review")
    incidencias = review.incidencias if review else []

    st.session_state["resultado_pipeline"] = {
        "draft_es": draft_es,
        "draft_en": draft_en,
        "incidencias": incidencias,
        "elapsed": elapsed,
    }

if "resultado_pipeline" in st.session_state:
    datos = st.session_state["resultado_pipeline"]
    draft_es = datos["draft_es"]
    draft_en = datos["draft_en"]
    incidencias = datos["incidencias"]
    elapsed = datos["elapsed"]

    # 3. Mostrar resultado
    st.markdown("---")
    st.markdown('<p class="zona-titulo">📝 Informe generado</p>', unsafe_allow_html=True)

    badge = (
        '<span class="badge-ok">✔ Validado</span>' if not incidencias
        else f'<span class="badge-warn">⚠ {len(incidencias)} incidencia(s)</span>'
    )
    st.markdown(f"""
    <div class="section-card">
        <h3>Informe (Español) &nbsp; {badge}</h3>
        <p style="color:#555;font-size:.82rem">⏱ {elapsed:.1f}s</p>
        <p>{draft_es.replace(chr(10), "<br>")}</p>
    </div>
    """, unsafe_allow_html=True)

    if draft_en:
        st.markdown(f"""
        <div class="section-card">
            <h3>Report (English)</h3>
            <p>{draft_en.replace(chr(10), "<br>")}</p>
        </div>
        """, unsafe_allow_html=True)

    if incidencias:
        with st.expander(f"⚠️ {len(incidencias)} incidencia(s) para revisión humana"):
            for inc in incidencias:
                st.warning(inc)

    # 4. Descarga
    st.divider()
    st.markdown('<p class="zona-titulo">💾 Descargar resultado</p>', unsafe_allow_html=True)
    docx_buffer = _construir_docx(draft_es, draft_en, incidencias)
    col_dl1, col_dl2, _ = st.columns([1, 1, 2])
    with col_dl1:
        st.download_button(
            label="📥 Descargar .docx",
            data=docx_buffer,
            file_name="memoria_anual_sanse.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            type="primary",
        )
    with col_dl2:
        texto_plano = f"INFORME (ESPAÑOL)\n\n{draft_es}\n\n\nREPORT (ENGLISH)\n\n{draft_en}"
        st.download_button(
            label="📥 Descargar .txt",
            data=texto_plano.encode("utf-8"),
            file_name="memoria_anual_sanse.txt",
            mime="text/plain",
        )
