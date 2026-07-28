"""report_generator.py - Generador del informe final.

Sintetiza la Memoria Anual siguiendo la estructura fija del ejemplo real
del Ayuntamiento (Introducción, Análisis por Plan, Actividad Operativa,
Conclusiones) a partir de state["analysis"] -- NO del draft libre del
Redactor.

La traducción al inglés REUTILIZA agente_adaptador_en, llamado aquí como función directa -- más
robusto que una implementación propia: reintento con fallback a texto
libre, limpieza de <think>, y compatible con el cambio de modelo por
cuota (adaptador_en.py ya sabe quitar reasoning_effort si el modelo no
es qwen). No se toca adaptador_en.py.

agente_adaptador_en se llama 4 veces (una por sección) desde
traducir_memoria() como función Python normal, NO como nodo de
LangGraph -- un nodo se ejecuta una vez por turno del grafo, y aquí
hace falta repetir la llamada con un texto distinto cada vez.

Trocea los datos de síntesis si no caben en una sola llamada (mismo
patrón que analyst.py y redactor_v1.py).

PENDIENTE:
- El documento puede salir en más páginas de las esperadas pese a
  respetar el presupuesto de palabras.
- No probado con los 3 documentos reales juntos en un solo run limpio.
- Sin plantilla oficial del Ayuntamiento.
"""

import os
import re
import unicodedata
from pathlib import Path

from docx import Document
from langchain_groq import ChatGroq
from langchain_core.prompts import ChatPromptTemplate

from src_agents.agents.adaptador_en import agente_adaptador_en
from src_agents.models.state import Analisis, ConceptoValor, EstadoPipeline

NOMBRE_ENTIDAD = os.environ.get("NOMBRE_ENTIDAD", "Ayuntamiento")

_MAX_BLOQUE_DATOS_CHARS = 2000

_PATRON_FUGA_RAZONAMIENTO = re.compile(
    r"\n?\d+\.\s*\*{0,2}(Analyze|Process|Draft|Task|Identify|Extract)\b",
    re.IGNORECASE,
)

EJEMPLO_ESTILO = """
Destaca la ampliación de las redes colaborativas, alcanzando las 78
personas integrantes de la Red, muy por encima de la meta prevista de
10. Esta colaboración se ha materializado en la ejecución de 3
proyectos en 2025 (surgidos de 0 en 2024), incluyendo ferias y
jornadas, con un nivel de satisfacción del 85%. La atención a personas
usuarias ha crecido exponencialmente, de 10 en 2024 a 508 en 2025.
"""

PROMPT_MEMORIA = """Eres un redactor técnico municipal. A partir de los
datos proporcionados, redacta una Memoria Anual de Actividades siguiendo
ESTRICTAMENTE este estilo (ejemplo de referencia, tono y densidad, no
copies su contenido):

"{ejemplo_estilo}"

Reglas:
- Usa EXCLUSIVAMENTE las cifras que aparecen en los datos proporcionados.
- No inventes cifras. Si un plan no tiene datos suficientes, dilo brevemente.
- Redacta SIEMPRE en párrafos fluidos, sin viñetas ni listas, en ninguna
  sección, integrando las cifras de forma natural en el texto.
- LÍMITE ESTRICTO: el conjunto de las 4 secciones no debe superar
  {presupuesto_palabras} palabras en total.

Responde EXACTAMENTE en este formato, sin nada más:
==INTRODUCCION==
(texto de la introducción)
==ANALISIS_PLANES==
(texto del análisis por plan)
==ACTIVIDAD_OPERATIVA==
(texto de la actividad operativa)
==CONCLUSIONES==
(texto de las conclusiones)

Datos disponibles:
{datos}
"""

_plantilla_memoria = ChatPromptTemplate.from_messages([("human", PROMPT_MEMORIA)])


def _limpiar_fuga_razonamiento(texto: str) -> str:
    """Corta el texto en el primer indicio de fuga de razonamiento sin
    <think> -- red de seguridad extra sobre lo que ya filtra
    adaptador_en.py, que solo limpia <think>...</think>."""
    match = _PATRON_FUGA_RAZONAMIENTO.search(texto)
    if match:
        return texto[:match.start()].strip()
    return texto


# --------------------------------------------------------------------
# Preparación de datos
# --------------------------------------------------------------------

def _deduplicar(analisis: Analisis) -> Analisis:
    """Agrupa ConceptoValor por valor único -- el mismo texto largo
    suele repetirse bajo varios conceptos distintos en tablas reales."""
    vistos = set()
    datos_dedup = []
    for d in analisis.datos:
        if d.valor not in vistos:
            vistos.add(d.valor)
            datos_dedup.append(d)
    return Analisis(datos=datos_dedup, notas=analisis.notas)


def _formatear_datos(analisis: Analisis) -> str:
    return "\n".join(f"- {d.concepto}: {d.valor}" for d in analisis.datos)


def _trocear_datos_texto(datos_texto: str, max_chars: int = _MAX_BLOQUE_DATOS_CHARS) -> list[str]:
    """Trocea el texto de datos en bloques que quepan en una sola
    llamada, sin cortar una línea de concepto a la mitad."""
    lineas = datos_texto.split("\n")
    bloques: list[str] = []
    bloque_actual: list[str] = []
    longitud_actual = 0
    for linea in lineas:
        if bloque_actual and longitud_actual + len(linea) > max_chars:
            bloques.append("\n".join(bloque_actual))
            bloque_actual = []
            longitud_actual = 0
        bloque_actual.append(linea)
        longitud_actual += len(linea) + 1
    if bloque_actual:
        bloques.append("\n".join(bloque_actual))
    return bloques


def _parsear_memoria(texto: str) -> dict:
    patron = r"==(\w+)==\s*(.*?)(?=\n==\w+==|\Z)"
    return {nombre: contenido.strip() for nombre, contenido in re.findall(patron, texto, re.DOTALL)}


def _quitar_acentos(texto: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", texto) if unicodedata.category(c) != "Mn")


# --------------------------------------------------------------------
# Validación
# --------------------------------------------------------------------

def validar_memoria(secciones: dict, analisis: Analisis, longitud_max: int = 15) -> list[str]:
    """Comprueba que los valores CORTOS (candidatos a cifra puntual) del
    Analista aparezcan en el texto generado. Ignora valores largos
    (nombres de indicador, párrafos narrativos)."""
    texto_normalizado = _quitar_acentos(" ".join(secciones.values()).lower())
    incidencias = []
    for dato in analisis.datos:
        if len(dato.valor) > longitud_max:
            continue
        if _quitar_acentos(dato.valor.lower()) not in texto_normalizado:
            incidencias.append(f"'{dato.concepto}' = {dato.valor} no aparece en la memoria generada")
    return incidencias


# --------------------------------------------------------------------
# Síntesis (con troceo)
# --------------------------------------------------------------------
_CLAVES_ESPERADAS = {"INTRODUCCION", "ANALISIS_PLANES", "ACTIVIDAD_OPERATIVA", "CONCLUSIONES"}


def _sintetizar_bloque(modelo, datos_texto_bloque: str, presupuesto_palabras: int) -> dict:
    prompt = _plantilla_memoria.invoke({
        "ejemplo_estilo": EJEMPLO_ESTILO,
        "presupuesto_palabras": presupuesto_palabras,
        "datos": datos_texto_bloque,
    })
    ultimo_error = None
    ultimo_resultado_parcial = {}
    for intento in range(3):
        try:
            respuesta = modelo.invoke(prompt)
            resultado = _parsear_memoria(respuesta.content)
            if not _CLAVES_ESPERADAS.issubset(resultado.keys()):
                faltantes = _CLAVES_ESPERADAS - resultado.keys()
                print(f"  [generador][debug] intento {intento+1}: faltan {faltantes}")
                print(f"  [generador][debug] tamaño del bloque de datos: {len(datos_texto_bloque)} caracteres")
                print(f"  [generador][debug] respuesta cruda:\n{respuesta.content[:800]}\n---")
                ultimo_error = ValueError(f"Faltan secciones en la respuesta: {faltantes}")
                ultimo_resultado_parcial = resultado
                continue
            return resultado
        except Exception as e:
            ultimo_error = e

    print(f"  [generador][aviso] bloque no completó las 4 secciones tras 3 intentos ({ultimo_error}); se usa el resultado parcial")
    return ultimo_resultado_parcial


def sintetizar_memoria(analisis: Analisis, presupuesto_palabras: int = 600) -> dict:
    """Llama a Groq para redactar las 4 secciones. Si no caben en una
    sola llamada, se trocean y se fusionan las secciones parciales."""
    modelo = ChatGroq(
        model=os.environ.get("GROQ_MODEL_MEMORIA", "openai/gpt-oss-120b"),
        api_key=os.environ["GROQ_API_KEY"],
        temperature=0.3,
        max_tokens=3000,
    )
    analisis_dedup = _deduplicar(analisis)
    datos_texto = _formatear_datos(analisis_dedup)
    bloques_datos = _trocear_datos_texto(datos_texto)

    if len(bloques_datos) == 1:
        return _sintetizar_bloque(modelo, bloques_datos[0], presupuesto_palabras)

    print(f"[generador] datos troceados en {len(bloques_datos)} bloques...")
    presupuesto_por_bloque = max(150, presupuesto_palabras // len(bloques_datos))
    secciones_parciales = []
    for indice, bloque in enumerate(bloques_datos, start=1):
        print(f"[generador] sintetizando bloque {indice}/{len(bloques_datos)}...")
        secciones_parciales.append(_sintetizar_bloque(modelo, bloque, presupuesto_por_bloque))

    claves = ["INTRODUCCION", "ANALISIS_PLANES", "ACTIVIDAD_OPERATIVA", "CONCLUSIONES"]
    return {
        clave: " ".join(s.get(clave, "").strip() for s in secciones_parciales if s.get(clave, "").strip())
        for clave in claves
    }


# --------------------------------------------------------------------
# Traducción — reutiliza el Adaptador del equipo, como función directa
# --------------------------------------------------------------------

def traducir_memoria(secciones: dict) -> dict:
    """Traduce las 4 secciones reutilizando agente_adaptador_en. Se
    llama 4 veces, una por sección -- no es un nodo del grafo, es una
    función que este módulo invoca directamente."""
    secciones_en = {}
    for clave, texto in secciones.items():
        resultado = agente_adaptador_en({"draft": texto})
        secciones_en[clave] = _limpiar_fuga_razonamiento(resultado.get("draft_en", ""))
    return secciones_en


# --------------------------------------------------------------------
# Documento final
# --------------------------------------------------------------------

def _anadir_secciones(doc: Document, secciones: dict, titulos: dict):
    for clave, titulo in titulos.items():
        doc.add_heading(titulo, level=1)
        doc.add_paragraph(secciones.get(clave, ""))


def generar_memoria_docx(analisis: Analisis, ruta_salida: Path, incluir_ingles: bool = True) -> Path:
    """Punto de entrada directo: sintetiza, traduce (reutilizando el
    Adaptador) y genera el .docx final, con las incidencias de
    validación como sección de notas."""
    secciones = sintetizar_memoria(analisis)
    incidencias = validar_memoria(secciones, _deduplicar(analisis))

    titulos_es = {
        "INTRODUCCION": "1. Introducción",
        "ANALISIS_PLANES": "2. Análisis de los Avances por Plan Estratégico",
        "ACTIVIDAD_OPERATIVA": "3. Análisis de la Actividad Operativa",
        "CONCLUSIONES": "4. Desenlace y Conclusiones",
    }

    doc = Document()
    doc.add_heading("Memoria Anual de Actividades 2025", level=0)
    doc.add_heading(f"Departamento de Innovación y Empleo — {NOMBRE_ENTIDAD}", level=2)
    _anadir_secciones(doc, secciones, titulos_es)

    if incluir_ingles:
        secciones_en = traducir_memoria(secciones)
        titulos_en = {
            "INTRODUCCION": "1. Introduction",
            "ANALISIS_PLANES": "2. Analysis of Progress by Strategic Plan",
            "ACTIVIDAD_OPERATIVA": "3. Analysis of Operational Activity",
            "CONCLUSIONES": "4. Outcome and Conclusions",
        }
        doc.add_page_break()
        doc.add_heading("Annual Activity Report 2025 (English version)", level=0)
        _anadir_secciones(doc, secciones_en, titulos_en)

    if incidencias:
        doc.add_page_break()
        doc.add_heading("Notas de validación (revisión humana)", level=1)
        p = doc.add_paragraph(
            "Datos puntuales del Analista no citados literalmente en el "
            "resumen (selección editorial esperable; revisar si algún "
            "dato clave falta):"
        )
        p.runs[0].italic = True
        for i in incidencias:
            doc.add_paragraph(i, style="List Bullet")

    doc.save(ruta_salida)
    return ruta_salida


def agente_generador_informe(estado: EstadoPipeline) -> dict:
    """Nodo de LangGraph: genera el .docx final a partir de
    state['analysis'] y devuelve la ruta en state['final_document'].
    Internamente llama a agente_adaptador_en como función (no como
    nodo separado del grafo) para la traducción."""
    ruta_salida = Path(os.environ.get("OUTPUT_DIR", "data/output")) / "memoria_final.docx"
    ruta_salida.parent.mkdir(parents=True, exist_ok=True)
    ruta = generar_memoria_docx(estado["analysis"], ruta_salida)
    return {"final_document": str(ruta)}