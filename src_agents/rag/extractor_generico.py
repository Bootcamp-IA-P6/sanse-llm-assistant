"""extractor_generico.py - Ingesta generica de documentos."""

import sys
from pathlib import Path
sys.path.append(str(Path(__file__).resolve().parents[2]))

from dataclasses import dataclass, field
from datetime import datetime
from typing import Literal, Union

ValorCelda = Union[str, int, float, None]

@dataclass
class BloqueContenido:
    """La unidad minima de informacion extraida de un documento."""
    tipo_bloque: Literal["texto", "tabla"]
    contenido: Union[str, list[list[ValorCelda]]]
    etiqueta: str
    fuente: str
    formato_origen: str
    fecha_extraccion: str = field(default_factory=lambda: datetime.now().isoformat())

# --------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------

import pdfplumber

def extraer_pdf(ruta: Path) -> list[BloqueContenido]:
    """Extrae texto y tablas de un PDF, pagina por pagina."""
    bloques = []
    with pdfplumber.open(ruta) as pdf:
        for num_pagina, pagina in enumerate(pdf.pages, start=1):
            tablas_detectadas = pagina.find_tables()
            pagina_sin_tablas = pagina
            for tabla in tablas_detectadas:
                pagina_sin_tablas = pagina_sin_tablas.outside_bbox(tabla.bbox)
            texto = pagina_sin_tablas.extract_text() or ""
            if texto.strip():
                bloques.append(BloqueContenido(
                    tipo_bloque="texto",
                    contenido=texto,
                    etiqueta=f"{ruta.stem} - pagina {num_pagina}",
                    fuente=ruta.name,
                    formato_origen="pdf",
                ))
            for tabla in tablas_detectadas:
                filas = tabla.extract()
                if filas:
                    bloques.append(BloqueContenido(
                        tipo_bloque="tabla",
                        contenido=filas,
                        etiqueta=f"{ruta.stem} - pagina {num_pagina}",
                        fuente=ruta.name,
                        formato_origen="pdf",
                    ))
    return bloques

# --------------------------------------------------------------------
# Word
# --------------------------------------------------------------------

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

def _iter_block_items(doc):
    """Recorre el documento en su orden real, mezclando parrafos y tablas."""
    for child in doc.element.body.iterchildren():
        if child.tag.endswith('}p'):
            yield Paragraph(child, doc)
        elif child.tag.endswith('}tbl'):
            yield Table(child, doc)

def _es_titulo(parrafo) -> bool:
    """Heuristica: un parrafo corto y en negrita actua como titulo."""
    if not parrafo.text.strip() or not parrafo.runs:
        return False
    return bool(parrafo.runs[0].bold) and len(parrafo.text) < 100

def _forward_fill_columna_agrupadora(tabla_filas: list[list[str]], indice_columna: int = 0) -> list[list[str]]:
    """Rellena celdas vacias de UNA columna concreta con el ultimo valor no
    vacio visto por encima, para reconstruir la columna de agrupacion
    (ej. MES) cuando una tabla de Word la deja en blanco por continuidad
    en vez de repetir el valor, como si hiciera Excel con una fusion real.
    No toca ninguna otra columna: una celda vacia en cualquier otra
    posicion es un dato ausente genuino, no una continuacion, y se deja
    tal cual (verificado contra filas TOTAL y datos faltantes reales)."""
    ultimo_valor = None
    filas_corregidas = []
    for fila in tabla_filas:
        fila = list(fila)
        valor_actual = fila[indice_columna].strip() if fila[indice_columna] else ""
        if valor_actual:
            ultimo_valor = valor_actual
        elif ultimo_valor is not None:
            fila[indice_columna] = ultimo_valor
        filas_corregidas.append(fila)
    return filas_corregidas

def extraer_docx(ruta: Path) -> list[BloqueContenido]:
    """Extrae parrafos y tablas de un Word, en su orden real."""
    doc = Document(ruta)
    bloques = []
    etiqueta_actual = ruta.stem
    buffer_texto = []

    def volcar_buffer():
        if buffer_texto:
            bloques.append(BloqueContenido(
                tipo_bloque="texto",
                contenido="\n".join(buffer_texto),
                etiqueta=etiqueta_actual,
                fuente=ruta.name,
                formato_origen="docx",
            ))
            buffer_texto.clear()

    for item in _iter_block_items(doc):
        if isinstance(item, Paragraph):
            if not item.text.strip():
                continue
            if _es_titulo(item):
                volcar_buffer()
                etiqueta_actual = item.text.strip().rstrip(":")
            else:
                buffer_texto.append(item.text)
        else:  # Table
            volcar_buffer()
            filas = [[celda.text for celda in fila.cells] for fila in item.rows]
            filas = _forward_fill_columna_agrupadora(filas, indice_columna=0)
            bloques.append(BloqueContenido(
                tipo_bloque="tabla",
                contenido=filas,
                etiqueta=etiqueta_actual,
                fuente=ruta.name,
                formato_origen="docx",
            ))

    volcar_buffer()
    return bloques

# --------------------------------------------------------------------
# Excel
# --------------------------------------------------------------------

from openpyxl import load_workbook

def _propagar_celdas_fusionadas(ws, filas: list) -> list:
    """Copia el valor de cada celda fusionada a TODAS las celdas que ocupa
    visualmente. openpyxl solo guarda el valor en la celda superior-
    izquierda del rango fusionado; el resto llegan vacias (None), lo que
    rompe la relacion entre una cabecera de grupo (ej. un titulo que cubre
    varias columnas) y las columnas que describe."""
    for rango in ws.merged_cells.ranges:
        valor = ws.cell(row=rango.min_row, column=rango.min_col).value
        for r in range(rango.min_row, rango.max_row + 1):
            idx_fila = r - ws.min_row
            if not (0 <= idx_fila < len(filas)):
                continue
            for c in range(rango.min_col, rango.max_col + 1):
                idx_col = c - ws.min_column
                if 0 <= idx_col < len(filas[idx_fila]):
                    filas[idx_fila][idx_col] = valor
    return filas

def extraer_xlsx(ruta: Path) -> list[BloqueContenido]:
    """Extrae cada hoja de un Excel como un bloque de tipo 'tabla', en
    formato de matriz cruda (sin asumir cual fila es la cabecera).

    NOTA: se carga SIN read_only=True porque ese modo no da acceso a
    ws.merged_cells - y sin esa informacion no se puede reconstruir
    correctamente una cabecera de grupo (un titulo que fusiona varias
    columnas). El coste en memoria es asumible para el tamano de
    archivo actual."""
    wb = load_workbook(ruta, data_only=True)
    bloques = []
    for nombre_hoja in wb.sheetnames:
        ws = wb[nombre_hoja]
        filas = [list(fila) for fila in ws.iter_rows(values_only=True)]
        filas = _propagar_celdas_fusionadas(ws, filas)
        filas = [f for f in filas if any(c is not None for c in f)]
        if not filas:
            continue
        bloques.append(BloqueContenido(
            tipo_bloque="tabla",
            contenido=filas,
            etiqueta=nombre_hoja,
            fuente=ruta.name,
            formato_origen="xlsx",
        ))
    return bloques

# --------------------------------------------------------------------
# Despachador
# --------------------------------------------------------------------

EXTRACTORES = {
    ".pdf": extraer_pdf,
    ".docx": extraer_docx,
    ".xlsx": extraer_xlsx,
}

def extraer_documento(ruta: Path) -> list[BloqueContenido]:
    """Punto de entrada unico: detecta el formato y llama al extractor."""
    extension = ruta.suffix.lower()
    if extension not in EXTRACTORES:
        raise ValueError(f"Formato no soportado: {extension} ({ruta.name})")
    return EXTRACTORES[extension](ruta)

def extraer_carpeta(carpeta: Path) -> list[BloqueContenido]:
    """Extrae todos los documentos soportados de una carpeta."""
    bloques = []
    for ruta in sorted(carpeta.iterdir()):
        if ruta.suffix.lower() in EXTRACTORES:
            print(f"Procesando: {ruta.name}")
            bloques.extend(extraer_documento(ruta))
    return bloques